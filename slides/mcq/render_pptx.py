# -*- coding: utf-8 -*-
"""Render the four new slides into the PPTX from browser-measured geometry."""
import json, shutil, html as _html
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

PX=6350
def E(px): return Emu(int(round(px*PX)))
def P(px): return Pt(round(px*0.5,1))
def C(h):
    h=h.lstrip("#"); return RGBColor(int(h[0:2],16),int(h[2:4],16),int(h[4:6],16))

LAY=json.load(open("layout.json"))
SRC="/Users/hadez/Documents/Company/training content/Agent_on_Databricks/Building Agentic AI Systems on Databricks.pptx"
BAK="pptx-before-new-slides.pptx"
shutil.copy2(SRC,BAK); print("backup ->",BAK)
prs=Presentation(SRC)
lst=prs.slides._sldIdLst
print("start:",len(lst))

ALIGN={"left":PP_ALIGN.LEFT,"center":PP_ALIGN.CENTER,"right":PP_ALIGN.RIGHT,
       "start":PP_ALIGN.LEFT,"end":PP_ALIGN.RIGHT}

def draw(sl, spec):
    sl.background.fill.solid(); sl.background.fill.fore_color.rgb=C("#FBFBF8")
    for b in spec["boxes"]:
        if b["kind"]=="arrow":
            s=sl.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,E(b["x"]),E(b["y"]),E(b["w"]),E(b["h"]))
            s.fill.solid(); s.fill.fore_color.rgb=C(b["fill"] or "#8B0000")
            s.line.fill.background(); s.shadow.inherit=False
            continue
        if not b.get("fill") and not b.get("line"): continue
        shape = MSO_SHAPE.ROUNDED_RECTANGLE if (b.get("radius") or 0)>=4 else MSO_SHAPE.RECTANGLE
        s=sl.shapes.add_shape(shape,E(b["x"]),E(b["y"]),E(b["w"]),E(b["h"]))
        if b.get("fill"): s.fill.solid(); s.fill.fore_color.rgb=C(b["fill"])
        else: s.fill.background()
        if b.get("line") and (b.get("lw") or 0)>0:
            s.line.color.rgb=C(b["line"]); s.line.width=Pt(max(0.5,(b.get("lw") or 1)*0.5))
        else:
            s.line.fill.background()
        s.shadow.inherit=False
        if (b.get("radius") or 0)>=4:
            try: s.adjustments[0]=min(0.5, b["radius"]/min(b["w"],b["h"]))
            except Exception: pass
        if b.get("leftbar"):
            lb=b["leftbar"]
            t=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,E(b["x"]),E(b["y"]),E(lb["w"]),E(b["h"]))
            t.fill.solid(); t.fill.fore_color.rgb=C(lb["c"] or "#8B0000")
            t.line.fill.background(); t.shadow.inherit=False
    for t in spec["texts"]:
        pad=2
        tb=sl.shapes.add_textbox(E(t["x"]-pad),E(t["y"]-pad),E(t["w"]+2*pad),E(t["h"]+2*pad))
        tf=tb.text_frame; tf.word_wrap=True
        tf.margin_left=tf.margin_right=E(pad); tf.margin_top=tf.margin_bottom=0
        p=tf.paragraphs[0]
        p.alignment=ALIGN.get(t.get("align","left"),PP_ALIGN.LEFT)
        try: p.line_spacing=max(0.8, t["lh"]/t["size"])
        except Exception: pass
        for r in t["runs"]:
            txt=r["t"]
            if t.get("upper"): txt=txt.upper()
            run=p.add_run(); run.text=txt
            f=run.font
            f.size=P(t["size"]); f.bold=bool(r["b"]); f.italic=bool(r["i"])
            f.color.rgb=C(t["color"] or "#14181F")
            f.name="JetBrains Mono" if r["m"] else "IBM Plex Sans"

def slide_of(el): return prs.part.related_part(el.rId).slide

# add_slide appends its sldId to the end of the list, so grab it right away
made={}
for spec in LAY:
    sl=prs.slides.add_slide(prs.slide_layouts[0])
    for ph in list(sl.placeholders): ph._element.getparent().remove(ph._element)
    draw(sl,spec)
    made[spec["sid"]]=list(lst)[-1]
    print(f"  drew {spec['sid']}: {len(spec['boxes'])} boxes, {len(spec['texts'])} text blocks")

# Insert by position in deck.json rather than by matching slide text: titles are not
# unique (the agenda lists every session name, so "Databricks Genie" matches slide 2).
DECK="/Users/hadez/Documents/Company/training content/Agent_on_Databricks/slides/project/deck.json"
order=json.load(open(DECK))["order"]
for el in made.values(): lst.remove(el)
assert len(lst)==len(order)-len(made), (len(lst), len(order), len(made))
for sid in sorted(made, key=order.index):
    lst.insert(order.index(sid), made[sid])
    print(f"  placed {sid} at deck position {order.index(sid)+1}")

print("slides now:",len(lst))

# renumber every page badge
deck=json.load(open("/Users/hadez/Documents/Company/training content/Agent_on_Databricks/slides/project/deck.json"))
renum=0
for n,el in enumerate(list(lst),1):
    sl=slide_of(el)
    for sh in sl.shapes:
        if not sh.has_text_frame: continue
        if sh.text_frame.text.strip().isdigit() and sh.left and sh.left>E(1690) and sh.top>E(940):
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.text.strip().isdigit():
                        if r.text!=str(n): renum+=1
                        r.text=str(n)
            break
print("badges renumbered:",renum)
prs.save(SRC)
print("saved ->",len(Presentation(SRC).slides._sldIdLst),"slides")
