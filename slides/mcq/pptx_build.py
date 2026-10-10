# -*- coding: utf-8 -*-
"""Apply the capstone removal + MCQ slides to the PPTX, matching the HTML deck."""
import copy, json, re, sys, shutil, html as _html
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
sys.path.insert(0, ".")
from questions import QUIZ

TARGET={"s1":[0,1,2,3],"s2":[1,2,3,0],"s3":[2,3,0,1],"s4":[3,0,1,2],
        "s5":[0,2,1,3],"s6":[1,3,0,2],"s7":[2,0,3,1]}
for _sid,(_t,_items) in QUIZ.items():
    new=[]
    for _i,(stem,opts,ci,why,oth) in enumerate(_items):
        tgt=TARGET[_sid][_i]
        rest=[o for j,o in enumerate(opts) if j!=ci]
        out=[None]*4; out[tgt]=opts[ci]
        it=iter(rest)
        for j in range(4):
            if out[j] is None: out[j]=next(it)
        new.append((stem,out,tgt,why,oth))
    QUIZ[_sid]=(_t,new)

PX=6350                      # EMU per px
def E(px): return Emu(int(round(px*PX)))
def P(px): return Pt(px*0.5) # 1920px == 960pt

INK=RGBColor(0x14,0x18,0x1F); SEC=RGBColor(0x4A,0x55,0x68)
MUT=RGBColor(0x8A,0x94,0xA3); RED=RGBColor(0x8B,0x00,0x00)
BG =RGBColor(0xFB,0xFB,0xF8); WHT=RGBColor(0xFF,0xFF,0xFF)
LIN=RGBColor(0xE3,0xE6,0xEA); RULE=RGBColor(0xD9,0xDC,0xE1)
LETTERS="ABCD"

def strip(s):
    s=re.sub(r"<[^>]+>","",s)
    return _html.unescape(s)
def mono_parts(s):
    """split on <span ...monospace...>..</span> so code keeps JetBrains Mono"""
    out=[];last=0
    for m in re.finditer(r"<span style='font-family:JetBrains Mono,monospace'>(.*?)</span>", s):
        if m.start()>last: out.append((strip(s[last:m.start()]),False))
        out.append((strip(m.group(1)),True)); last=m.end()
    if last<len(s): out.append((strip(s[last:]),False))
    return [(t,m) for t,m in out if t]

def txbox(sl,x,y,w,h):
    tb=sl.shapes.add_textbox(E(x),E(y),E(w),E(h))
    tf=tb.text_frame; tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    return tf

def runs(par, parts, size, color, bold=False):
    for t,m in parts:
        r=par.add_run(); r.text=t
        r.font.size=P(size); r.font.color.rgb=color; r.font.bold=bold
        r.font.name="JetBrains Mono" if m else "IBM Plex Sans"

def rect(sl,x,y,w,h,fill,line=None,shape=MSO_SHAPE.ROUNDED_RECTANGLE,adj=None):
    s=sl.shapes.add_shape(shape,E(x),E(y),E(w),E(h))
    s.fill.solid(); s.fill.fore_color.rgb=fill
    if line is None: s.line.fill.background()
    else: s.line.color.rgb=line; s.line.width=Pt(0.75)
    s.shadow.inherit=False
    if adj is not None:
        try: s.adjustments[0]=adj
        except Exception: pass
    return s

def chrome(sl, pagenum):
    sl.background.fill.solid(); sl.background.fill.fore_color.rgb=BG
    rect(sl,1908,0,12,1080,RED,shape=MSO_SHAPE.RECTANGLE)          # right bar
    rect(sl,128,1008,1664,1,RULE,shape=MSO_SHAPE.RECTANGLE)        # footer rule
    rect(sl,128,1004,56,6,INK,shape=MSO_SHAPE.RECTANGLE)           # tick
    b=rect(sl,1716,1012,76,38,RED,shape=MSO_SHAPE.RECTANGLE)
    tf=b.text_frame; tf.word_wrap=False
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
    r=p.add_run(); r.text=str(pagenum); r.font.size=P(24); r.font.bold=True
    r.font.color.rgb=RGBColor(0xFB,0xFB,0xF8); r.font.name="IBM Plex Sans"

def header(sl,title,tag,tagcol,lead):
    tf=txbox(sl,128,80,1664,64); p=tf.paragraphs[0]
    r=p.add_run(); r.text=title; r.font.size=P(48); r.font.bold=True
    r.font.color.rgb=INK; r.font.name="IBM Plex Sans"
    r=p.add_run(); r.text="     "+tag; r.font.size=P(24); r.font.bold=True
    r.font.color.rgb=tagcol; r.font.name="IBM Plex Sans"
    tf=txbox(sl,128,152,1664,32); p=tf.paragraphs[0]
    runs(p, mono_parts(lead), 23, SEC)

ROW_Y0, ROW_H, ROW_GAP = 188, 194, 12
def row_box(sl,i):
    y=ROW_Y0+i*(ROW_H+ROW_GAP)
    rect(sl,128,y,1664,ROW_H,WHT,LIN,adj=0.045)
    return y

def badge(sl,y,n):
    c=rect(sl,156,y+22,38,38,RED,shape=MSO_SHAPE.OVAL)
    tf=c.text_frame; tf.word_wrap=False
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
    r=p.add_run(); r.text=str(n); r.font.size=P(21); r.font.bold=True
    r.font.color.rgb=RGBColor(0xFB,0xFB,0xF8); r.font.name="IBM Plex Sans"

def question_slide(prs, title, items, pagenum):
    sl=prs.slides.add_slide(prs.slide_layouts[0])
    for ph in list(sl.placeholders): ph._element.getparent().remove(ph._element)
    chrome(sl,pagenum)
    header(sl,title,"Knowledge check",MUT,
           "Four questions. One correct answer each — answers on the next slide.")
    for i,(stem,opts,_,_,_) in enumerate(items):
        y=row_box(sl,i); badge(sl,y,i+1)
        tf=txbox(sl,205,y+20,540,ROW_H-36)
        runs(tf.paragraphs[0], mono_parts(stem), 22, INK, bold=True)
        tf=txbox(sl,786,y+20,978,ROW_H-36)
        for j,o in enumerate(opts):
            p=tf.paragraphs[0] if j==0 else tf.add_paragraph()
            p.space_after=P(9)
            r=p.add_run(); r.text=LETTERS[j]+"   "
            r.font.size=P(20); r.font.bold=True; r.font.color.rgb=INK; r.font.name="IBM Plex Sans"
            runs(p, mono_parts(o), 20, SEC)
    return sl

def answer_slide(prs, title, items, key, pagenum):
    sl=prs.slides.add_slide(prs.slide_layouts[0])
    for ph in list(sl.placeholders): ph._element.getparent().remove(ph._element)
    chrome(sl,pagenum)
    header(sl,title,"Answers",RED, key+"  — and why the others fail.")
    for i,(stem,opts,ci,why,others) in enumerate(items):
        y=row_box(sl,i); badge(sl,y,i+1)
        tf=txbox(sl,205,y+20,540,ROW_H-36)
        runs(tf.paragraphs[0], [(LETTERS[ci]+" — ",False)]+mono_parts(opts[ci]), 22, INK, bold=True)
        tf=txbox(sl,786,y+18,978,ROW_H-32)
        p=tf.paragraphs[0]; p.space_after=P(8); runs(p, mono_parts(why), 20, SEC)
        p=tf.add_paragraph(); runs(p, mono_parts(others), 19, MUT)
    return sl

# ---------------------------------------------------------------- main
SRC="/Users/hadez/Documents/Company/training content/Agent_on_Databricks/Building Agentic AI Systems on Databricks.pptx"
BAK="/private/tmp/claude-501/-Users-hadez-projects-datacouch-moodle/15e19203-e54b-41c6-b187-203637057ea4/scratchpad/pptx-backup.pptx"
DECK="/Users/hadez/Documents/Company/training content/Agent_on_Databricks/slides/project/deck.json"
shutil.copy2(SRC,BAK); print("backup ->", BAK)

prs=Presentation(SRC)
lst=prs.slides._sldIdLst
ids=list(lst)
print("start:", len(ids))

def text_of(sl):
    return " ".join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame)

# 2. build the 14 quiz slides and move them into place
order=json.load(open(DECK))["order"]
pos={sid:n for n,sid in enumerate(order,1)}
made={}
for sid,(title,items) in QUIZ.items():
    key=", ".join(f"{n}{LETTERS[items[n-1][2]]}" for n in range(1,5))
    q=question_slide(prs,title,items,pos[f"{sid}-quiz"])
    a=answer_slide(prs,title,items,key,pos[f"{sid}-quiz-answers"])
    made[sid]=(q,a,key)
    print(f"  built {sid}: pages {pos[f'{sid}-quiz']}/{pos[f'{sid}-quiz-answers']}  key {key}")

# 1. drop the three capstone slides
drop=[]
for i,sl in enumerate(prs.slides):
    t=text_of(sl)[:90].lower()
    if ("capstone · 4 hours" in t or "the rubric, in the open" in t
        or "capstone · the grader" in t):
        drop.append(i)
print("dropping capstone slides at 1-based:", [i+1 for i in drop])
assert len(drop)==3, drop
for i in sorted(drop, reverse=True):
    el=ids[i]
    prs.part.drop_rel(el.rId)
    lst.remove(el)
print("after capstone removal:", len(lst))

LAST={"s1":"s1-t5-best","s2":"s2-t5-why","s3":"s3-t6-why","s4":"s4-bench",
      "s5":"s5-t5-best","s6":"s6-t4-action","s7":"s7-t6-pitfalls"}
# locate anchors by their on-slide title text
ANCHOR_TEXT={
 "s1":"Context, memory and injection",
 "s2":"Why framework choice matters less",
 "s3":"Why the best agent is often no agent",
 "s4":"Benchmarks measure it. The agent inherits it.",
 "s5":"Tools, errors and approval",
 "s6":"See it in action: three aliases, one lever",
 "s7":"The pitfalls this course actually hit"}

ids=list(lst)
def slide_of(el):
    return prs.part.related_part(el.rId).slide

ANCHOR_TEXT={
 "s1":"Context, memory and injection",
 "s2":"Why framework choice matters less",
 "s3":"Why the best agent is often no agent",
 "s4":"Benchmarks measure it. The agent inherits it.",
 "s5":"Tools, errors and approval",
 "s6":"See it in action: three aliases, one lever",
 "s7":"The pitfalls this course actually hit"}

for sid in ["s1","s2","s3","s4","s5","s6","s7"]:
    q,a,_=made[sid]
    cur=list(lst)
    qid=[e for e in cur if slide_of(e) is q][0]
    aid=[e for e in cur if slide_of(e) is a][0]
    anchor=None
    for e in cur:
        if e in (qid,aid): continue
        t=" ".join(sh.text_frame.text for sh in slide_of(e).shapes if sh.has_text_frame)
        if ANCHOR_TEXT[sid] in t: anchor=e; break
    assert anchor is not None, sid
    lst.remove(qid); lst.remove(aid)
    j=list(lst).index(anchor)
    lst.insert(j+1,qid); lst.insert(j+2,aid)
    print(f"  placed {sid} after {ANCHOR_TEXT[sid][:34]!r}")

print("after insert:", len(lst))

# 3. renumber every page badge to its new position
NO={1}  # cover
renum=0
for n,el in enumerate(list(lst),1):
    sl=prs.part.related_part(el.rId).slide
    for sh in sl.shapes:
        if not sh.has_text_frame: continue
        t=sh.text_frame.text.strip()
        if t.isdigit() and sh.left is not None and sh.left>E(1690) and sh.top>E(940):
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.text.strip().isdigit():
                        if r.text!=str(n): renum+=1
                        r.text=str(n)
            break
print("page badges renumbered:", renum)
prs.save(SRC)
print("saved:", SRC, "->", len(Presentation(SRC).slides._sldIdLst), "slides")
